"""Rebuild general.docx from general.md content — ATS-friendly, matching brand formatting."""
from docx import Document
from docx.shared import Pt, Emu, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Calibri"
COLOR_BORDER = "2F4B4F"
COLOR_LINK = RGBColor(0x2F, 0x4B, 0x4F)

NAME_SIZE = Pt(20)
CONTACT_SIZE = Pt(10)
SECTION_SIZE = Pt(12)
ROLE_SIZE = Pt(10.5)
BODY_SIZE = Pt(10)
SUMMARY_SIZE = Pt(10)

SP_NONE = Pt(0)
SP_1 = Emu(12700)
SP_2 = Emu(25400)
SP_3 = Emu(38100)
SP_4 = Emu(50800)
SP_6 = Emu(76200)
SP_8 = Emu(101600)
SP_10 = Emu(127000)

LINE_SPACING_TWIPS = 260  # 13pt in twips (twentieths of a point) for w:line


def set_font(run, size=BODY_SIZE, bold=False, color=None):
    run.font.name = FONT
    run.font.size = size
    run.bold = bold
    run.font.color.rgb = color if color else RGBColor(0, 0, 0)
    # Ensure Calibri is used for East Asian fallback too (ATS compat)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), FONT)
    rFonts.set(qn("w:cs"), FONT)


def set_line_spacing(paragraph, spacing_twips=LINE_SPACING_TWIPS):
    """Set explicit line spacing. spacing_twips is in twips (20ths of a point)."""
    pPr = paragraph._element.get_or_add_pPr()
    sp = pPr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        pPr.append(sp)
    sp.set(qn("w:line"), str(spacing_twips))
    sp.set(qn("w:lineRule"), "atLeast")


def add_bottom_border(paragraph):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), COLOR_BORDER)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_hyperlink(paragraph, url, text):
    """Add a proper hyperlink that ATS parsers can follow."""
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run_elem = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    rStyle = OxmlElement("w:rStyle")
    rStyle.set(qn("w:val"), "Hyperlink")

    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)
    rFonts.set(qn("w:eastAsia"), FONT)
    rFonts.set(qn("w:cs"), FONT)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(CONTACT_SIZE.pt * 2)))  # half-points

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2F4B4F")

    # No underline — cleaner look
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "none")

    rPr.append(rStyle)
    rPr.append(rFonts)
    rPr.append(sz)
    rPr.append(color)
    rPr.append(u)
    run_elem.append(rPr)

    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run_elem.append(t)

    hyperlink.append(run_elem)
    paragraph._element.append(hyperlink)


def add_section_header(doc, text):
    """Use Heading 2 style for ATS section recognition, then apply brand formatting."""
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    p.paragraph_format.space_before = SP_10
    p.paragraph_format.space_after = SP_4
    p.clear()
    run = p.add_run(text)
    set_font(run, SECTION_SIZE, bold=True)
    add_bottom_border(p)
    set_line_spacing(p)
    return p


def add_role(doc, title, location):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = SP_6
    p.paragraph_format.space_after = SP_NONE
    run = p.add_run(title)
    set_font(run, ROLE_SIZE, bold=True)
    set_line_spacing(p)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = SP_NONE
    p2.paragraph_format.space_after = SP_1
    run2 = p2.add_run(location)
    set_font(run2, BODY_SIZE)
    run2.italic = True
    set_line_spacing(p2)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = SP_NONE
    p.paragraph_format.space_after = SP_1
    p.clear()
    run = p.add_run(text)
    set_font(run, BODY_SIZE)
    set_line_spacing(p)


def add_skill_line(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = SP_NONE
    p.paragraph_format.space_after = SP_1
    run_label = p.add_run(f"{label}: ")
    set_font(run_label, BODY_SIZE, bold=True)
    run_value = p.add_run(value)
    set_font(run_value, BODY_SIZE)
    set_line_spacing(p)


def add_project_header(doc, title, tech):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = SP_6
    p.paragraph_format.space_after = SP_NONE
    run_title = p.add_run(title)
    set_font(run_title, ROLE_SIZE, bold=True)
    run_tech = p.add_run(f", {tech}")
    set_font(run_tech, BODY_SIZE)
    run_tech.italic = True
    set_line_spacing(p)


def build():
    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.page_width = 7772400
    section.page_height = 10058400
    section.left_margin = 548640
    section.right_margin = 548640
    section.top_margin = 457200
    section.bottom_margin = 457200

    # Set Normal style defaults
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = BODY_SIZE
    normal.font.color.rgb = RGBColor(0, 0, 0)

    # Set List Bullet style defaults
    lb = doc.styles["List Bullet"]
    lb.font.name = FONT
    lb.font.size = BODY_SIZE
    lb.font.color.rgb = RGBColor(0, 0, 0)

    # Configure Heading 2 style to not look like default Word heading
    h2 = doc.styles["Heading 2"]
    h2.font.name = FONT
    h2.font.size = SECTION_SIZE
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.font.all_caps = False

    # === NAME (Heading 1 for ATS) ===
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = SP_NONE
    p.paragraph_format.space_after = SP_2
    p.clear()
    run = p.add_run("Barry Babbitt")
    set_font(run, NAME_SIZE, bold=True)

    # Configure Heading 1 style
    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT
    h1.font.color.rgb = RGBColor(0, 0, 0)

    # === CONTACT with hyperlinks ===
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = SP_NONE
    p.paragraph_format.space_after = SP_3

    r = p.add_run("Orlando, FL  |  ")
    set_font(r, CONTACT_SIZE)

    add_hyperlink(p, "mailto:babbitt.barry@gmail.com", "babbitt.barry@gmail.com")
    r = p.add_run("  |  ")
    set_font(r, CONTACT_SIZE)

    r = p.add_run("407-451-3979")
    set_font(r, CONTACT_SIZE)
    r = p.add_run("  |  ")
    set_font(r, CONTACT_SIZE)

    add_hyperlink(p, "https://linkedin.com/in/barrybabbitt1", "linkedin.com/in/barrybabbitt1")
    r = p.add_run("  |  ")
    set_font(r, CONTACT_SIZE)

    add_hyperlink(p, "https://bbabbitt1.github.io/portfolio", "bbabbitt1.github.io/portfolio")

    # === PROFESSIONAL SUMMARY ===
    add_section_header(doc, "PROFESSIONAL SUMMARY")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = SP_NONE
    p.paragraph_format.space_after = SP_3
    run = p.add_run(
        "Analytics and data engineer who designs dimensional models and builds the "
        "pipelines and apps that run on them. Ships production work on Azure, AWS, and "
        "Snowflake, from API ingestion through Power BI semantic models and full-stack web apps."
    )
    set_font(run, SUMMARY_SIZE)
    set_line_spacing(p)

    # === SKILLS ===
    add_section_header(doc, "SKILLS")

    skills = [
        ("Programming", "SQL, Python, TypeScript, React, Node.js, M, DAX, C#, HTML, CSS, JavaScript"),
        ("Cloud & Warehousing", "AWS (S3, Lambda, DMS, Kinesis, ECR), Azure (Functions, SQL), Snowflake, BigQuery"),
        ("Data Engineering", "ETL/ELT pipelines, schema evolution, SCD1/SCD2, streaming data, data lakehouse, CI/CD (GitLab/GitHub)"),
        ("Application Development", "React, Express, JWT authentication, PWA (Workbox, IndexedDB), REST API design, Zod validation"),
        ("Analytics & BI", "Power BI, Power Query, Power Automate, Excel, Azure Data Studio, dbt, Erwin Data Modeler"),
        ("Data Architecture", "Star schemas, dimensional modeling, conceptual/logical/physical models, grain alignment"),
        ("SQL Server (Advanced)", "Relational modeling, CTE-heavy transformations, window functions, stored procedures, views"),
        ("API Integration", "REST API consumption, deeply nested JSON flattening (Python, C#), SQLAlchemy"),
        ("Frameworks", "Agile/Scrum, A/B Testing, canonical schemas (ISO 20022)"),
    ]
    for label, value in skills:
        add_skill_line(doc, label, value)

    # === EXPERIENCE ===
    add_section_header(doc, "EXPERIENCE")

    add_role(doc,
        "Senior Business Analyst, Disney (Analytics Suite Engineering)",
        "Orlando, FL | September 2026 \u2013 Present")
    add_bullet(doc, "Building a full-stack web app that brings the internal analytics suite into one product for business users.")
    add_bullet(doc, "Turning recurring BI and reporting requests into self-service features in the app.")

    add_role(doc,
        "Senior Data Analyst, Transitions (Reporting Automation & Data Governance)",
        "Orlando, FL | May 2026 \u2013 September 2026")
    add_bullet(doc, "Led a full redesign of the Power BI semantic model, improving data quality and cutting refresh time from 9s to 3s.")
    add_bullet(doc, "Built and owned a Python reporting platform that runs 9+ manual data pulls from one command, with validation, logging, and a report-freshness dashboard. It replaced a multi-hour weekly routine.")
    add_bullet(doc, "Extracted a 115-tab, 22MB trended P&L workbook into a self-reconciling budget fact table, checking each line item against the workbook\u2019s subtotals and ending monthly manual re-keying.")
    add_bullet(doc, "Automated per-site report distribution for 11 site leaders via the Power BI REST API and generated PowerPoint decks, cutting hours of manual assembly each pay period.")
    add_bullet(doc, "Set up data governance: per-table data dictionaries with a controlled key vocabulary, a type/length canon across warehouses for spotting drift, and a published metric glossary.")

    add_role(doc,
        "Senior Data Analyst, BNY (Enterprise Data Modeling & Governance)",
        "Lake Mary, FL | March 2025 \u2013 April 2026")
    add_bullet(doc, "Architected a metadata and data-dictionary platform with GitLab and SharePoint API integrations, documenting 15+ enterprise datasets and 2,000+ attributes in one place.")
    add_bullet(doc, "Built a GitLab CI/CD pipeline that validates, versions, and promotes documentation through environments, cutting schema-change turnaround from days to hours.")
    add_bullet(doc, "Acted as the team\u2019s de facto technical lead: clarified tasks, set best practices, and helped teammates through implementation, which sped up onboarding.")

    add_role(doc,
        "Contract Data Engineer",
        "Orlando, FL | 2025")
    add_bullet(doc, "Engineered a production data pipeline using Tekla PowerFab APIs (C#), Python ETL, and AWS S3, automating ingestion of fabrication metrics into SQL Server and delivering a Power BI dashboard leadership uses to track daily production.")
    add_bullet(doc, "Implemented Type 2 SCD tracking using dbt and Snowflake across 2M+ records, so the business could audit compliance and see how attributes changed over time.")

    add_role(doc,
        "Senior Data Analyst, MasterCorp",
        "Orlando, FL | October 2023 \u2013 March 2025")
    add_bullet(doc, "Introduced and owned star-schema semantic models across finance, operations, HR, and legal using surrogate keys and effective dating (SCD2), giving the company its first consistent cross-department reporting.")
    add_bullet(doc, "Designed the company\u2019s first enterprise HR analytics dashboard in Power BI and SQL and defined turnover metrics across siloed departments. Leadership used it in a cross-functional initiative that cut turnover 25%.")
    add_bullet(doc, "Built an Operations Financial Overview dashboard using advanced DAX and SQL, evaluating 200+ sites on financial, operational, and merit-based KPIs with automated bonus-calculation logic.")

    add_role(doc,
        "Financial Planning Analyst, Raymond James",
        "Orlando, FL | January 2022 \u2013 October 2023")
    add_bullet(doc, "Automated capital gains exposure monitoring using BlackRock\u2019s API integrated with SQL-driven holdings data, identifying and flagging $50M in avoidable distributions and shaping distribution strategy.")
    add_bullet(doc, "Delivered 600+ regulatory-aligned financial plans for 100+ financial advisors, using SQL to check client holdings against compliance standards.")

    add_role(doc,
        "Reporting Analyst, Universal Property & Casualty Insurance",
        "Ft. Lauderdale, FL | January 2021 \u2013 January 2022")
    add_bullet(doc, "Replaced manual regulatory reporting for claims and underwriting with SQL stored procedures, giving finance and compliance faster, more accurate reports.")

    add_role(doc,
        "Data Analyst & Assistant Coach, Rollins College",
        "Orlando, FL | August 2018 \u2013 January 2022")
    add_bullet(doc, "Built a SQL database centralizing player performance, TrackMan telemetry, and tournament data, and optimized practice schedules with linear programming, which raised attendance 50%.")

    # === PROJECTS ===
    add_section_header(doc, "PROJECTS")

    add_project_header(doc, "Shots Gained Tracker", "React 19, TypeScript, Express, Azure SQL, PWA, IndexedDB")
    add_bullet(doc, "Built a mobile-first PWA in React 19, TypeScript, and Express on Azure SQL for live strokes gained tracking against ~3,500 PGA Tour benchmarks, with running totals by category (Driving, Approach, Short Game, Putting).")
    add_bullet(doc, "Designed an offline-first architecture with IndexedDB and Workbox service workers: cached benchmarks let you log rounds and calculate SG without signal, and a sync queue uploads them when you reconnect.")

    add_project_header(doc, "Blast Motion Putting Pipeline", "Python, Azure Functions, Azure SQL, Power Automate, Power BI")
    add_bullet(doc, "Engineered a serverless ETL pipeline using Azure Functions (timer + HTTP triggers) and Python, extracting putting stroke data from Blast Motion\u2019s REST API and loading it into Azure SQL with a high-watermark pattern and idempotent upserts.")
    add_bullet(doc, "Built a one-tap Power Automate workflow that runs the pipeline, refreshes the dataset, and sends a push notification, getting new putts onto the dashboard in under 60 seconds.")

    # === EDUCATION ===
    add_section_header(doc, "EDUCATION")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = SP_NONE
    p.paragraph_format.space_after = SP_2
    run = p.add_run("Bachelor of Arts in Mathematics")
    set_font(run, BODY_SIZE, bold=True)
    run2 = p.add_run(", Rollins College, Orlando, FL (May 2018)")
    set_font(run2, BODY_SIZE)
    set_line_spacing(p)

    # === CERTIFICATIONS ===
    add_section_header(doc, "CERTIFICATIONS")
    certs = [
        "Microsoft DP-600: Fabric Analytics Engineer",
        "AWS Solutions Architect Associate",
        "Microsoft PL-300: Power BI Data Analyst Associate",
        "Microsoft AZ-900: Azure Cloud Fundamentals",
        "Google Advanced Data Analytics Professional Certificate",
    ]
    for cert in certs:
        add_bullet(doc, cert)

    # Save
    output = "C:/Users/babbi/development/portfolio/resumes/general.docx"
    doc.save(output)
    print(f"Saved: {output}")


if __name__ == "__main__":
    build()
